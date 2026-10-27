# -*- coding: utf-8 -*-
"""
生成现场路演 PPT：提交材料/雨证-路演PPT.pptx
用法：python 04-脚本/make-ppt.py [输出路径]

注意（2026-10-12）：本脚本产出的是 **26 页对照件**（`提交材料/雨证-路演PPT.pptx`），
它已于 2026-10-11 移出仓库、只在需要重建时跑；**现行提交母版是另一份手工版
`提交材料/神麻组-雨证-路演PPT1.pptx`（12 页），不由本脚本生成**。

叙事版：乔布斯式的节奏（一页一个想法、大字、留白、悬念→揭晓）+ 雷军式的对比与参数页。
每页的演讲者备注 = 6 分钟口播（含停顿提示），PowerPoint 演讲者视图可直接用。

文字与数字来自仓库里已核实的源头：提交材料/演示讲稿.md、
02-作战与答辩/汉客松-交易哈希清单.md、10-金融与定价/指标推导-变量表.md。
改数字前先改这些源头。
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- 设计 token（与 DESIGN.md 一致）
BG = RGBColor(0x07, 0x0E, 0x19)
PANEL = RGBColor(0x0C, 0x17, 0x25)
PANEL2 = RGBColor(0x10, 0x1E, 0x30)
INK = RGBColor(0xE9, 0xF0, 0xF8)
DIM = RGBColor(0xA9, 0xBA, 0xCB)
FAINT = RGBColor(0x7F, 0x93, 0xA9)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
BLUEL = RGBColor(0x7F, 0xB0, 0xFF)
ACCENT = RGBColor(0xFF, 0xC8, 0x57)
OK = RGBColor(0x3F, 0xD9, 0xA5)
BAD = RGBColor(0xF0, 0x83, 0x6F)
HAIR = RGBColor(0x22, 0x33, 0x4A)
RAINLINE = RGBColor(0x12, 0x20, 0x3A)

FONT = "微软雅黑"
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.9)
CW = W - 2 * M

SHOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "08-截图存证", "演示页-968-视口.png")
ASSETS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                      "08-截图存证", "ppt-assets")


def _set_cjk(run, name=FONT):
    rPr = run.font._rPr
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def picture(slide, name, x, y, w):
    p = os.path.join(ASSETS, name)
    if not os.path.exists(p):
        return None
    return slide.shapes.add_picture(p, x, y, width=w)


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, first=False, before=0, after=0, line=1.2):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.line_spacing = line
    return p


def run(p, text, size=14, color=INK, bold=False, space=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    _set_cjk(r)
    if space is not None:
        r.font._rPr.set("spc", str(int(space * 100)))
    return r


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def hairline(slide, x, y, w, color=HAIR, weight=0.75):
    return rect(slide, x, y, w, Pt(weight), color)


def blank(prs, rain=0):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    for i in range(rain):                     # 极低对比雨丝
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2 + i * 0.62), Inches(-0.4), Pt(0.75), Inches(8.2))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RAINLINE
        sh.line.fill.background()
        sh.shadow.inherit = False
        sh.rotation = 12
    return s


def foot(slide, n):
    tf = textbox(slide, M, H - Inches(0.52), CW, Inches(0.24))
    p = para(tf, True)
    run(p, "雨证 · RainProof", 9, FAINT)
    run(p, "      队伍 神麻 · 汉客松 S1 & ETH Wuhan 2026", 9, FAINT)
    tf2 = textbox(slide, W - M - Inches(0.8), H - Inches(0.52), Inches(0.8), Inches(0.24), PP_ALIGN.RIGHT)
    run(para(tf2, True), f"{n:02d}", 9, FAINT, True)


def big(slide, text, y, size=40, color=INK, bold=True, x=M, w=CW, align=PP_ALIGN.LEFT, line=1.25):
    tf = textbox(slide, x, y, w, Inches(2.0), align)
    for i, t in enumerate(text if isinstance(text, list) else [text]):
        run(para(tf, i == 0, before=0 if i == 0 else 6, line=line), t, size, color, bold)
    return tf


def kicker(slide, text, y=Inches(0.62), color=FAINT):
    tf = textbox(slide, M, y, CW, Inches(0.26))
    run(para(tf, True), text, 10.5, color, True, space=1.6)


def bullets(slide, x, y, w, items, size=14, gap=10, h=Inches(3.6), line=1.35):
    tf = textbox(slide, x, y, w, h)
    for i, (lead, rest) in enumerate(items):
        p = para(tf, i == 0, before=0 if i == 0 else gap, line=line)
        if lead:
            run(p, lead, size, ACCENT, True)
        if rest:
            run(p, rest, size, DIM)
    return tf


def metrics(slide, y, cells, x=M, w=CW, h=Inches(1.35), num_size=30, lab_size=10.5):
    n = len(cells)
    gap = Inches(0.2)
    cw = Emu(int((w - gap * (n - 1)) / n))
    for i, (num, lab, color) in enumerate(cells):
        cx = Emu(int(x + i * (cw + gap)))
        rect(slide, cx, y, cw, h, PANEL, HAIR)
        tf = textbox(slide, cx + Inches(0.2), y + Inches(0.22), cw - Inches(0.4), h - Inches(0.4))
        run(para(tf, True), num, num_size, color, True)
        run(para(tf, before=5, line=1.25), lab, lab_size, DIM)


def table(slide, x, y, w, rows, col_w, row_h=Inches(0.4), head_size=11, body_size=11.5):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Emu(int(row_h * len(rows))))
    tbl = shp.table
    for j, cwd in enumerate(col_w):
        tbl.columns[j].width = cwd
    for i, row in enumerate(rows):
        tbl.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL2 if i == 0 else PANEL
            cell.margin_left = cell.margin_right = Inches(0.13)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(head_size if i == 0 else body_size)
            r.font.bold = i == 0 or j == 1
            r.font.color.rgb = BLUEL if i == 0 else (INK if j == 1 else DIM)
            r.font.name = FONT
            _set_cjk(r)
    return tbl


def flow(slide, y, steps, h=Inches(1.3)):
    n = len(steps)
    gap = Inches(0.36)
    cw = Emu(int((CW - gap * (n - 1)) / n))
    for i, (title, desc) in enumerate(steps):
        cx = Emu(int(M + i * (cw + gap)))
        rect(slide, cx, y, cw, h, PANEL, HAIR)
        tf = textbox(slide, cx + Inches(0.16), y + Inches(0.16), cw - Inches(0.32), h - Inches(0.28))
        run(para(tf, True), title, 14, INK, True)
        run(para(tf, before=5, line=1.2), desc, 10, FAINT)
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(int(cx + cw + Inches(0.07))), y + Inches(0.5), Inches(0.22), Inches(0.2))
            ar.fill.solid()
            ar.fill.fore_color.rgb = BLUE
            ar.line.fill.background()
            ar.shadow.inherit = False


def small(slide, text, y, size=10.5, color=FAINT, x=M, w=CW):
    tf = textbox(slide, x, y, w, Inches(0.6))
    run(para(tf, True, line=1.3), text, size, color)
    return tf


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- 故事线
def s01_cover(prs, n):
    s = blank(prs, rain=14)
    kicker(s, "汉客松 S1 & ETH Wuhan 2026 · BOT Chain 赛道", Inches(1.15))
    big(s, ["如果一场雨，", "能让一个骑手白干一天"], Inches(1.75), size=40, line=1.3)
    tf = textbox(s, M, Inches(3.9), Inches(9.5), Inches(0.9))
    run(para(tf, True, line=1.35), "那他不该只能认命。", 19, ACCENT, True)
    run(para(tf, before=8), "雨证 · RainProof —— 可验证的 AI 参数化配送险", 13.5, DIM)
    run(para(tf, before=6), "把「你信不信我」换成「你自己验」", 13.5, ACCENT, True)
    hairline(s, M, Inches(5.0), Inches(3.2), BLUE, 2)
    tf = textbox(s, M, Inches(5.5), CW, Inches(1.0))
    run(para(tf, True, line=1.5), "https://juice-butterfly.github.io/rainproof/", 12.5, BLUEL)
    run(para(tf, line=1.5), "github.com/juice-butterfly/rainproof     队伍 神麻", 11, FAINT)
    notes(s, "【开场，别急】各位评委好，我们是神麻。今天我不打算先讲条款和技术，我想先讲一个人 —— 一个很具体的人。")
    return s


def s02_person(prs, n):
    s = blank(prs, rain=8)
    kicker(s, "10 月 8 日 · 武汉 · 凌晨五点")
    big(s, ["他还是出门了。"], Inches(1.55), size=34)
    tf = textbox(s, M, Inches(2.9), Inches(10.4), Inches(2.6))
    for i, t in enumerate([
        (True, "暴雨橙色预警。他穿上雨衣，骑上车，跑了 12 个小时。"),
        (False, "当天不出勤，就没有收入。"),
        (False, "这一场雨，让他今天的收入少了三成 —— 而他不知道该找谁赔。"),
    ]):
        p = para(tf, i == 0, before=0 if i == 0 else 14, line=1.4)
        run(p, t[1], 18 if not t[0] else 19, INK if t[0] else DIM, t[0])
    small(s, "外卖骑手是千万级的群体，而「暴雨断收入」这件事，几乎不在任何保单的保障范围里。", Inches(5.6), size=12.5, color=FAINT)
    notes(s, "武汉，凌晨五点，暴雨橙色预警。他还是出门了 —— 因为当天不出勤，就没有收入。跑了一天，收入少了三成。"
             "他不知道该找谁赔。【停半秒】这不是一个故事，这是每天。")
    return s


def s03_conflict(prs, n):
    s = blank(prs)
    big(s, ["他不需要同情。", "他需要一份真的能赔的保险。"], Inches(2.3), size=33, line=1.32)
    hairline(s, M, Inches(4.6), Inches(3.2), BLUE, 2)
    small(s, "但市面上几乎没有这个险种 —— 而这个「几乎没有」，原因并不在条款里。", Inches(5.0), size=14, color=DIM)
    notes(s, "他不需要同情，他需要的是一份真的能赔的保险。但市面上几乎没有这个险种。为什么？原因不在条款里。")
    return s


def s04_math(prs, n):
    s = blank(prs)
    kicker(s, "一笔账")
    big(s, ["因为定损的成本，比赔款还高。"], Inches(1.5), size=30)
    metrics(s, Inches(2.85), [
        ("几十元", "一单赔款", ACCENT),
        ("同量级", "查勘一次的成本", BAD),
        ("以周计", "报案 → 查勘 → 核赔", DIM),
    ], h=Inches(1.5), num_size=26)
    small(s, "所以不是没人想保 —— 是这笔账算不过来。这一层不解决，条款写得再漂亮也落不了地。", Inches(4.85), size=14, color=DIM)
    notes(s, "一单赔几十块，但查勘一次的成本和赔款是同一个量级。再叠上报案、查勘、核赔，周期以周计。"
             "对一个「当天断了收入」的人，这笔账算不过来。")
    return s


def s05_what_if(prs, n):
    s = blank(prs)
    big(s, ["如果，我们根本不问他有没有损失呢？"], Inches(3.0), size=36, line=1.3)
    notes(s, "【停两秒，再开口】如果，我们根本不问他有没有损失呢？因为我们要保的那件事，本来就不需要他自己证明。")
    return s


def s06_idea(prs, n):
    s = blank(prs)
    kicker(s, "方案：参数化保险")
    big(s, ["只问一个问题：", "保单期内，这个区域下了多少雨。"], Inches(1.35), size=27, line=1.3)
    metrics(s, Inches(3.35), [
        ("50 mm", "触发阈值 · 写在合约里", BLUE),
        ("0.01 ETH", "每笔赔付 · 写在合约里", ACCENT),
        ("24 / 48 / 72 h", "保障期三档 · v2 才分档", OK),
    ], h=Inches(1.4), num_size=27)
    small(s, "阈值与赔付额是 constant —— 部署之后，连我们自己都改不了。保费是可调的承保价，每次改动都在链上留一条依据哈希。", Inches(5.3), size=12.5, color=DIM)
    notes(s, "我们的规则只有一句话：保单期内，这个区域累计降水的增量超过 50 毫米，赔付 0.01 ETH。"
             "阈值、赔付额写在合约里，是 constant —— 部署之后连我们自己都改不了。"
             "保费不一样，它是承保价，可调，但每次改动都要在链上留一条记录和依据哈希。")
    return s


def s07_flow(prs, n):
    s = blank(prs)
    kicker(s, "一次赔付，五步走完")
    if not picture(s, "链路五步-flow.png", M, Inches(1.45), CW):
        flow(s, Inches(1.5), [
            ("① 投保", "链上记下雨量快照"),
            ("② 喂价", "两层闸门，三模型交叉"),
            ("③ AI 判定", "R1–R5 规则，哈希上链"),
            ("④ 触发", "claim 权限开放"),
            ("⑤ 赔付", "只进保单登记地址"),
        ], h=Inches(1.5))
    bullets(s, M, Inches(4.95), Inches(6.6), [
        ("每一步都在链上留一条事件。", "PolicyBought / RainfallUpdated / JudgementSubmitted / ClaimPaid。"),
    ], size=14, h=Inches(1.3))
    rect(s, Inches(7.9), Inches(4.85), Inches(4.53), Inches(1.6), PANEL, HAIR)
    tf = textbox(s, Inches(8.2), Inches(5.07), Inches(3.9), Inches(1.3))
    run(para(tf, True), "一句必须说清楚的话", 10.5, FAINT, True, space=1.2)
    run(para(tf, before=7, line=1.35), "合约没有自动触发机制：claim 权限开放，任何人可触发，钱只进保单登记的骑手地址。我们只是替自己跑了个 keeper。", 12, DIM)
    notes(s, "投保、喂价、AI 判定、触发、赔付，五步。每一步都在链上留一条事件。"
             "这里要讲清楚一件事：合约本身没有「自动」触发机制 —— claim 权限是开放的，任何人可触发，"
             "钱只会进保单里登记的骑手地址。我们只是替自己跑了个 keeper 而已。")
    return s


def s08_demo(prs, n):
    s = blank(prs)
    kicker(s, "现在，用真实数据跑一遍")
    if os.path.exists(SHOT):
        s.shapes.add_picture(SHOT, M, Inches(1.35), width=Inches(7.6))
        rect(s, M, Inches(1.35), Inches(7.6), Inches(4.65), None, HAIR)
    bullets(s, Inches(8.85), Inches(1.45), Inches(3.6), [
        ("① 拨一下区域。", "武汉 → 广州：看这条刻度越过触发线。"),
        ("② 读保单 #0。", "上海、34mm、判定 DENY、置信度 95。"),
        ("③ 粘一个哈希。", "进核验台，直接向节点要原始记录。"),
    ], size=12.5, gap=12)
    small(s, "线上页面直连 BOT Chain 测试网 968 的节点，没有服务器、没有区块浏览器、没有缓存。", Inches(6.2), size=11)
    notes(s, "【切到演示页，30 秒】这是我们线上正在跑的页面，直连节点。我现在把区域从武汉拨到广州 —— "
             "看这条刻度，已经越过触发线。再看这份保单：#0，上海，34 毫米，AI 判定 DENY，置信度 95。"
             "这条保单不会赔一分钱。")
    return s


def s09_punchline(prs, n):
    s = blank(prs)
    big(s, ["这不是失败。", "这是它没有为了演示而撒谎。"], Inches(2.0), size=36, line=1.32, color=INK)
    hairline(s, M, Inches(4.5), Inches(3.2), ACCENT, 2)
    small(s, "保单 #0 · 上海 · 34mm → DENY · 置信度 95 · 交易 0x867ad825…c20ba0（真实链上数据，不是预录）", Inches(4.85), size=12, color=DIM)
    notes(s, "为什么我特意把这条「不赔」的保单端给你看？因为如果我们的系统今天愿意为了演示赔一次，"
             "你就有理由怀疑它明天会不会为了别的原因赔一次。【停一秒】它没有为了演示而撒谎。")
    return s


def s10_ai1(prs, n):
    s = blank(prs)
    kicker(s, "AI 第一件事 · 喂价门禁")
    big(s, ["让写进链上的那个数，值得信。"], Inches(1.35), size=27)
    bullets(s, M, Inches(2.75), Inches(7.0), [
        ("三个独立预报模型交叉核验。", "ECMWF / GFS / ICON 各算一遍；多数对不上，就不写进链。"),
        ("今天下午，它真的拒收过一次。", "广州三个模型 42.4 / 56.5 / 88.2mm，差 45.8mm ≫ 容差 11.3mm → 判 diverge，调 rejectFeed 留证。"),
        ("链上那个数，一动没动。", "累计值仍是 50mm —— 拒收本身也是一个上链动作。"),
    ], size=13)
    small(s, "交易 0x9e3f4b1c5fe9c0de74aa5dea420a45f99562bfd2d27cc2d0820679e6b6844a9f · 区块 11,860,989 · FeedRejected(region 4, confidence 40, sources 3)", Inches(5.9), size=10)
    notes(s, "写数的预言机是单一运营方，它写什么链上就信什么。所以 AI 的第一件事，是给这个数当门禁。"
             "同时取三个独立的气象模型交叉核验。今天下午它真的拒收过一次：广州三个模型差了 45 毫米，"
             "判定分歧，调 rejectFeed 留证 —— 链上那个累计值一动没动。闸门不是橡皮图章。")
    return s


def s11_ai2(prs, n):
    s = blank(prs)
    kicker(s, "AI 第二件事 · 判定复核")
    big(s, ["链上的数字，和三个模型背离超过 60%，", "就不赔。"], Inches(1.35), size=27, line=1.3)
    bullets(s, M, Inches(3.1), Inches(7.2), [
        ("R1–R5 是确定性规则，不是大模型随口说。", "最值钱的一条是 R2：偏离超过 60% 直接判 DENY。"),
        ("大模型只写人话解释，解释不进哈希。", "换一份快照，inputHash 必然变。"),
        ("同一份快照重跑，outputHash 逐字节一致。", "check-ai.js 会当场验给你看。"),
    ], size=13)
    picture(s, "三模型交叉-R1R5.png", Inches(8.55), Inches(3.4), Inches(3.9))
    small(s, "真实数据案例：上海 34mm → DENY 95；注入模拟的成都 120mm → PAY，模拟标记写进证据哈希。", Inches(5.85), size=11.5)
    notes(s, "第二件事是复核。规则是确定性的五条 R1–R5，最值钱的一条是 R2：链上官方数字和三个模型的中位数"
             "背离超过 60%，直接判不赔。大模型只写人话解释，而且解释不参与哈希 —— 同一份快照重跑，"
             "输出哈希逐字节一致。")
    return s


def s12_ai3(prs, n):
    s = blank(prs)
    kicker(s, "AI 第三件事 · 承保定价")
    big(s, ["11 年 · 五座城市 · 每小时的雨，", "算出一个价格。"], Inches(1.3), size=27, line=1.3)
    metrics(s, Inches(3.15), [
        ("10%", "保本线（10 倍杠杆）", DIM),
        ("2.27%", "北京 · 安全", OK),
        ("3.72%", "成都 · 安全", OK),
        ("4.76%", "武汉 · 安全", OK),
        ("11.57%", "广州 · 越线 → 加价", BAD),
    ], h=Inches(1.6), num_size=22, lab_size=10)
    small(s, "触发概率低于 10% 才保本；广州 95% 区间 9.35%–13.88%，所以链上单独加价（0.0020 RISK_LOADED）。这是精算给产品的定价权。", Inches(5.3), size=12)
    notes(s, "第三件事是定价。我们用 11 年、五座城市、逐小时的真实降雨算触发概率。"
             "10 倍杠杆的保本线是 10%：北京 2.27、成都 3.72、武汉 4.76，都在线内；广州 11.57%，越线 —— "
             "所以链上广州单独加了价。这就是精算带来的定价权。不要说「广州一定亏」，讲区间。")
    return s


def s13_onemore(prs, n):
    s = blank(prs)
    kicker(s, "One more thing", color=ACCENT)
    big(s, ["链上一落保单，AI 就自己做一次独立复核"], Inches(1.45), size=29)
    bullets(s, M, Inches(2.85), Inches(7.4), [
        ("PolicyBought 一触发，五项检查跑一遍。", "窗口是否在 24/48/72 内 · 基线单调 · 喂价新鲜度 ≤ 24h · 三模型是否认这份形势 · 与中位数背离 ≤ 60%。"),
        ("reviewHash = keccak256(canonical(review))。", "进哈希的链上值全部取自投保所在区块，喂价窗口写死 —— 换个扫描窗口重跑，哈希不变。"),
        ("它不碰链上状态。", "不改、不阻断、也不代替判定 —— 这是一条证据链，不是第四个执行者。"),
    ], size=12.5, h=Inches(2.4))
    metrics(s, Inches(5.5), [
        ("8 份", "两条链已留痕（Sepolia 2 / BOT 6）", ACCENT),
        ("19 项", "喂价闸门 + 承保复核断言", OK),
    ], h=Inches(1.0), num_size=22)
    notes(s, "One more thing。我们做了个钩子：链上一落保单，AI 就自动对这份保单做一次独立复核，"
             "五项检查，全部写成可复算的哈希。两条链已经留了 8 份。它不碰链上状态 —— 这是刻意的。")
    return s


def s14_compare(prs, n):
    s = blank(prs)
    kicker(s, "对比")
    big(s, ["传统方案 vs 雨证"], Inches(1.3), size=27)
    table(s, M, Inches(2.65), CW, [
        ["", "传统骑手天气险", "雨证 RainProof"],
        ["定损依据", "报案 + 查勘 + 核赔", "保单期内该区域降水增量"],
        ["定损成本", "与赔款同量级", "一次链上读，约 2.5e-4 ETH"],
        ["到账周期", "以周计", "判定上链之后当场"],
        ["谁来触发", "理赔员决定", "claim 权限开放，任何人可触发"],
        ["可核验性", "内部系统，事后无法复算", "事件、哈希、留痕全在链上，第三方可复算"],
        ["覆盖人群", "风控筛掉高风险城市", "参数化定价，高风险城市可保（加价）"],
    ], [Inches(1.5), Inches(4.6), Inches(5.33)], row_h=Inches(0.45), body_size=11)
    small(s, "「把定损成本降到接近零」，说的就是从第一行到第二行的这个变化。", Inches(6.15), size=11.5, color=FAINT)
    notes(s, "把两边摆在一起看。传统方案：定损靠报案和查勘，成本与赔款同量级，到账以周计，谁触发由理赔员决定。"
             "雨证：定损是链上读一个数，到账在判定上链之后，触发权限开放，全流程可核验。"
             "这就是「定损成本降到接近零」的意思。")
    return s


def s15_cost(prs, n):
    s = blank(prs)
    kicker(s, "但我们发现了更重要的一件事")
    big(s, ["决定价格的，不是风险，是成本。"], Inches(1.35), size=29)
    metrics(s, Inches(2.85), [
        ("≈ 2.5e-4 ETH", "每笔上链的固定成本", BAD),
        ("9.08e-6 ~ 1.236e-4", "分档后的公平保费（全部低于地板）", OK),
    ], h=Inches(1.55), num_size=22, lab_size=10.5)
    small(s, "五座城市 × 三档期限，风险折算出来的价格全都低于保费地板 —— 定价的瓶颈不是风险，而是那一笔固定成本。"
             "反过来说：链上的成本降一个量级，这个产品的价格就能降一个量级。", Inches(5.1), size=13, color=DIM)
    notes(s, "算完分档价格之后，我们发现一件更有意思的事：把风险折算成公平保费，五座城市、三档期限，"
             "全部低于我们设的保费地板。也就是说，定价的瓶颈不是风险，而是每笔上链的固定成本，大约 2.5e-4 ETH。"
             "这句话的另一面是好消息：链上成本降一个量级，这个产品就能便宜一个量级。")
    return s


def s16_spec(prs, n):
    s = blank(prs)
    kicker(s, "一页读完全部参数（全部读链，不是宣传材料）")
    table(s, M, Inches(1.45), CW, [
        ["参数", "值", "出处"],
        ["每笔赔付", "0.01 ETH", "constant，部署后不可改"],
        ["触发阈值", "50 mm / 24h", "constant；v2 按 GB/T 28592-2012 分三档"],
        ["保障期", "24 / 48 / 72 小时", "v2 A1，合约 require 限定"],
        ["保费", "0.0002（地板）~ 0.0020 ETH", "链上 premiumGrid + MIN_PREMIUM"],
        ["单骑手上限", "3 笔 / 0.02 ETH 在保敞口", "v2 A3"],
        ["准备金", "reserveOf() = max(reserve, openExposure)", "v2 A6，自动下限"],
        ["白名单", "默认关闭", "v2 A8"],
    ], [Inches(1.9), Inches(5.3), Inches(4.23)], row_h=Inches(0.44), body_size=11)
    small(s, "两条链：BOT Chain 测试网 968 是演示页读的那一条（v2 承保层，九项改动，彩排 25 项断言全绿）；"
             "Sepolia 11155111 上是第一阶段的 v1 留痕（真实数据 DENY 在这里）。地址同为 0x89e7C942535930B61cB61631051E8b0bD670596a —— "
             "同一个地址在两条链上是两个合约，绝不混写；数值单位在 Sepolia 上是 ETH、在 BOT Chain 上是 BOT（合约常量同一套）。", Inches(5.9), size=11.5, color=DIM)
    notes(s, "这一页是我们全部参数，全部从链上读出来，不是宣传材料。两条链分开说：演示页读的是 BOT Chain 968 上的 v2，"
             "Sepolia 上是第一阶段的 v1 留痕。地址一样，但是两个合约 —— 这一点我们从不混着讲。")
    return s


def s17_close(prs, n):
    s = blank(prs, rain=6)
    big(s, ["雨是数据。", "判定是代码。打款是自动的。"], Inches(2.1), size=33, line=1.32)
    hairline(s, M, Inches(4.35), Inches(3.2), ACCENT, 2)
    small(s, "规则和证据都在链上，你自己验。", Inches(4.7), size=15, color=DIM)
    tf = textbox(s, M, Inches(5.5), CW, Inches(1.2))
    for i, t in enumerate([
        "演示页  https://juice-butterfly.github.io/rainproof/",
        "源码    github.com/juice-butterfly/rainproof",
        "核验    单文件核验台 · 两条链 · 8 份可复算的 AI 留痕",
    ]):
        run(para(tf, i == 0, before=0 if i == 0 else 8, line=1.4), t, 12, BLUEL if i == 0 else DIM)
    picture(s, "演示站二维码.png", Inches(10.85), Inches(4.45), Inches(1.7))
    small(s, "扫码进演示站", Inches(6.25), size=10, color=FAINT, x=Inches(10.85), w=Inches(1.7))
    notes(s, "【收尾，慢一点】雨是数据，判定是代码，打款是自动的。规则和证据都在链上，你自己验。谢谢。"
             "注意：紧接着必须补一句 —— 合约没有自动触发机制，claim 权限开放、任何人可触发，"
             "钱只进保单登记的骑手地址。")
    return s


def s18_appendix(prs, n):
    s = blank(prs)
    kicker(s, "附录 · 可当场复算")
    big(s, ["你现在就能验的三件事"], Inches(1.3), size=25)
    bullets(s, M, Inches(2.5), Inches(5.9), [
        ("跑一遍断言。", "十七道门禁全绿。"),
        ("重算一份哈希。", "8 份 AI 判定留痕，输入与输出哈希可逐字节重算。"),
        ("自己问节点。", "两个页面都不是区块浏览器，没有缓存，直接向 RPC 要原始记录。"),
    ], size=12.5, gap=10)
    rect(s, Inches(7.15), Inches(2.5), Inches(5.28), Inches(3.3), PANEL, HAIR)
    tf = textbox(s, Inches(7.45), Inches(2.7), Inches(4.7), Inches(3.0))
    run(para(tf, True), "当场可复算", 10.5, FAINT, True, space=1.2)
    for t in ["cd 07-测试工具 && npm test —— 十七道门禁全绿（步数与断言数在 package.json 的 test 链里）",
              "node check-feed-verify.js —— 喂价闸门 12 项 + 承保复核 7 项",
              "node push-rainfall.js 4 —— 真跑一遍喂价，走闸门、必要时拒收",
              "09-AI判定留痕/ —— 8 份 JSON，哈希可当场重算"]:
        run(para(tf, before=9, line=1.3), t, 11, DIM)
    notes(s, "这一页留给问答：评委问到我们做不到的地方，答案就在这里。数字都可以当场复算。")
    return s


# ---------------------------------------------------------------- B 的 PPT 内容（合并进来）
def s04b_problem(prs, n):
    s = blank(prs)
    kicker(s, "四个卡点")
    big(s, ["为什么到今天，还是没人做这件事"], Inches(1.3), size=27)
    table(s, M, Inches(2.55), CW, [
        ["卡点", "具体表现", "雨证的回应"],
        ["定损成本 > 赔款", "一个骑手赔几十元，查勘一次是同量级", "不查勘：只读链上降水增量"],
        ["损失根本核不了", "没有平台订单数据，证明不了「少赚了」", "不核个人损失：区域参数化触发"],
        ["逆选择与骗赔", "看着要下雨才投保，事后补单", "冷静期 + 喂价新鲜度 ≤ 24h + 必须窗口前投保"],
        ["身份与凭据", "众包骑手与平台的关系不稳定", "v2 白名单（默认关闭）／B2B2C 走平台 attestation"],
    ], [Inches(2.35), Inches(5.2), Inches(3.98)], row_h=Inches(0.66), body_size=11)
    small(s, "四条里有三条是「成本与激励」问题，只有最后一条是「身份」问题 —— 所以我们先打成本这一层。",
          Inches(6.05), size=12, color=DIM)
    notes(s, "为什么没人做？第一，定损成本比赔款还高；第二，骑手的收入损失根本核不了；"
             "第三，知道要下雨才来买。这三条都是成本与激励结构问题。第四，身份和凭据不稳定 —— "
             "那要靠商务落地，所以我们先解决前三层。")
    return s


def s04c_target(prs, n):
    s = blank(prs)
    kicker(s, "为谁做的")
    big(s, ["平台众包骑手：按单计酬，不出勤就是零收入。"], Inches(1.3), size=25)
    metrics(s, Inches(2.5), [
        ("336 万", "月均有单骑手（美团研究院 2024）", ACCENT),
        ("6.8 倍", "雨天配送事故率（泰国骑手研究 OR = 6.79）", BAD),
        ("5 座", "首批城市：武汉 / 上海 / 北京 / 广州 / 成都", DIM),
    ], h=Inches(1.5), num_size=26)
    bullets(s, M, Inches(4.4), CW, [
        ("他们不是雇员。", "没有底薪，现有意外险只保人身伤害，不保收入中断。"),
        ("所以我们要保的不是「他受伤了」，而是「这场雨让他没法接单」。", "而这正好是一个只靠公开数据就能判定的问题。"),
    ], size=13, gap=12, h=Inches(1.9))
    notes(s, "为谁做的？平台上的众包骑手：336 万月均有单骑手，按单计酬，不出勤就是零收入。"
             "雨天配送的事故率大约是平时的 6.8 倍。但现有的意外险只保人身伤害，不保收入中断。"
             "所以我们要保的是「这场雨让他没法接单」—— 而这恰好是一个只靠公开数据就能判定的问题。")
    return s


def s07b_arch(prs, n):
    s = blank(prs)
    kicker(s, "技术方案与架构")
    big(s, ["五层架构：接缝是哈希，不是模型。"], Inches(1.15), size=24)
    picture(s, "架构五层.png", M, Inches(1.85), Inches(5.9))
    bullets(s, Inches(7.15), Inches(2.0), Inches(5.3), [
        ("数据来源单一化。", "链上只接受一个喂价入口，其余一律拒收。"),
        ("哈希是唯一的跨层契约。", "层与层之间只交换哈希，不交换文本。"),
        ("写链操作只有一个出口。", "keeper 轮询 + 三模型交叉，写链前先算 canonical 哈希。"),
        ("v2 的九项改动全部落在这一层。", "A1 保费网格 / A3 敞口上限 / A4 冷静期 / A4b 喂价新鲜度 / A5 来源数 / "
                                          "A6 敞口账本 / A7 国标分档 / A8 白名单 / A9 productId。"),
    ], size=11.5, gap=9, h=Inches(4.6))
    notes(s, "技术上分五层：外部数据、AI 判定、共识与结算、自动化与证据、产品前端。三条硬规则："
             "数据来源单一化；哈希是唯一的跨层契约；写链操作只有一个出口。一句话 —— 接缝是哈希，不是模型。")
    return s


def s09b_evidence(prs, n):
    s = blank(prs)
    kicker(s, "三笔证据 · 两条链")
    big(s, ["不赔的那一笔，比赔了的那一笔更值钱。"], Inches(1.25), size=26)
    table(s, M, Inches(2.45), CW, [
        ["#", "落在哪里", "数据与结果"],
        ["0", "Sepolia 真链 11155111", "上海 34mm（真实数据）→ DENY · 置信 95 · 一分不赔"],
        ["1", "Sepolia 真链 11155111", "成都 6 → 120mm（注入模拟暴雨，标记写进证据哈希）→ PAY 94 · 赔付 0.01 ETH"],
        ["—", "本地回放链（真实历史数据）", "武汉 2024-06-27 ~ 06-29 真实累计 230mm · simulated = false → PAY 76"],
    ], [Inches(0.6), Inches(3.2), Inches(7.73)], row_h=Inches(0.62), body_size=11)
    metrics(s, Inches(5.05), [
        ("≈ 0.00025 ETH", "含判定与赔付的链上总成本", OK),
        ("2.5%", "占 0.01 ETH 赔付额的比例", DIM),
    ], h=Inches(1.2), num_size=24)
    small(s, "第三笔不是链上交易：它跑在本地回放链上（chainId 伪装成 11155111），数据是 Open-Meteo 的真实历史数据 —— "
             "我们不会把它说成「链上赔付」。", Inches(6.42), size=10.5, color=FAINT)
    notes(s, "三笔证据。第一笔：上海 34 毫米，真实数据，判定不赔，置信 95 —— 这是我们的诚实基线。"
             "第二笔：成都，注入的模拟暴雨，赔付 0.01 ETH，模拟标记写进了证据哈希。"
             "第三笔要说清楚：它跑在本地回放链上，用的是武汉 2024 年 6 月那场真实暴雨的公开数据，230 毫米，"
             "判定赔付。我们不把它说成链上交易。")
    return s


def s14b_tiers(prs, n):
    s = blank(prs)
    kicker(s, "v2 承保层 · 按国标分档")
    big(s, ["赔付不再是一口价：按 GB/T 28592-2012 分三档。"], Inches(1.25), size=25)
    table(s, M, Inches(2.5), CW, [
        ["档位", "国标阈值（24h）", "赔付比例", "72h 窗口下的等效阈值", "11 年实测"],
        ["暴雨", "50 mm", "50%", "150 mm", "会触发 —— 演示与历史回放都在这一档"],
        ["大暴雨", "100 mm", "75%", "300 mm", "五城 11 年里几乎没有"],
        ["特大暴雨", "250 mm", "100%", "750 mm", "11 年未出现（武汉极值 392 mm）"],
    ], [Inches(1.4), Inches(1.9), Inches(1.2), Inches(2.6), Inches(4.43)], row_h=Inches(0.58), body_size=10.5)
    small(s, "门槛写死成 constant，敞口按档计入 openExposure；0 档赔 50%，即 0.005 ETH。"
             "目标赔付率 60%，余下 40% 覆盖成本与利润 —— 按此重定价，五城实际赔付率可以落到 53%–60%。",
          Inches(5.7), size=12, color=DIM)
    notes(s, "v2 里我们把「一口价」改成国标三档：暴雨 50 毫米赔一半，大暴雨赔四分之三，特大暴雨全赔。"
             "72 小时窗口下，二、三档要 300 和 750 毫米 —— 五城 11 年实测几乎不触发，所以演示和历史回放都落在第一档。"
             "门槛写死在合约里，敞口按档计入准备金。")
    return s


def s17b_business(prs, n):
    s = blank(prs)
    kicker(s, "商业模式")
    big(s, ["骑手自己不付钱。付钱的是有合规与舆情压力的一方。"], Inches(1.25), size=24)
    table(s, M, Inches(2.6), Inches(6.5), [
        ["谁付钱", "判断"],
        ["骑手自付", "不推荐：单价低，获客成本高"],
        ["平台 / 众包服务商", "首选：一次决策覆盖成千上万骑手"],
        ["保险公司", "第二阶段：成为它的参数化承保能力"],
    ], [Inches(2.3), Inches(4.2)], row_h=Inches(0.62), body_size=11)
    bullets(s, Inches(7.7), Inches(2.6), Inches(4.75), [
        ("区域分散。", "单城风险不相关，看组合口径而不是单城口径。"),
        ("动态定价。", "区域 × 时长 × 季节的三维网格。"),
        ("再保险。", "广州水文年 72h 触发频率 6.77%–16.10%，可以分层分出。"),
        ("扩品类。", "同一套骨架可移植到高温 ≥ 40℃、大风 17.2 m/s、空气质量。"),
    ], size=11.5, gap=9, h=Inches(3.4))
    notes(s, "商业模式一句话：骑手不付钱。付钱的是有合规压力和舆情压力的一方 —— 平台或者众包服务商。"
             "它做一次决策就能覆盖成千上万个骑手，边际获客成本接近零。保险公司是第二阶段。"
             "放大路径四条：区域分散、动态定价、再保险、扩品类 —— 后三个险种用的是同一套骨架。")
    return s


def s18a_roadmap(prs, n):
    s = blank(prs)
    kicker(s, "附录 · 后续规划")
    big(s, ["四阶段：从国标分档，到险种活化。"], Inches(1.15), size=24)
    picture(s, "路线图四阶段.png", M, Inches(1.85), Inches(5.9))
    bullets(s, Inches(7.15), Inches(2.0), Inches(5.3), [
        ("v2 九项已落地并部署。", "BOT Chain 测试网 968：0x89e7C942…596a —— 与 Sepolia 同址，但是两个合约。"),
        ("e2e_v2 101 项断言全绿。", "真链彩排 25 项全过：保费网格、敞口上限、国标分档、第三方 claim。"),
        ("02 / 03 / 04 还没排期。", "防逆选择（冷静期已实现，其余待做）、身份认定（需商务落地）、"
                                    "险种活化（验证后启动）。"),
    ], size=11.5, gap=9, h=Inches(4.6))
    notes(s, "后续四阶段：国标分档已经落地在 v2；防逆选择是下一步；身份认定需要商务落地；险种活化要等验证。"
             "v2 九项已经部署到 BOT Chain 测试网 968，e2e_v2 101 项断言全绿，真链彩排 25 项全过。")
    return s


def s18b_limits(prs, n):
    s = blank(prs)
    kicker(s, "附录 · 诚实边界（六条，每条都带处置）")
    table(s, M, Inches(1.35), CW, [
        ["边界", "现在是什么样", "怎么处置"],
        ["身份无校验", "没有 KYC，任何人都能买", "B2B2C + 平台 attestation 白名单（v2 A8，默认关闭）"],
        ["保费曾与时长无关", "v1 里 1h 与 72h 同一个价", "v2 A1：区域 × 时长 网格，只卖 24/48/72 + 地板"],
        ["预言机与判定是单点", "目前由单一运营方写数", "3-reporter 2/3 多签 → TEE / zkML"],
        ["v1 准备金没有自动规则", "「卫兵」已实现，但 reserve 手工设定，当前为 0", "v2 已自动化：reserveOf() = max(reserve, openExposure)"],
        ["赔付条件曾不联动", "v1 的 claim 不看有没有 AI 判定", "v2 已把「已有 PAY 判定」变成 claim 的前提"],
        ["L1 主网经济不成立", "北京档 0.0004 ETH < 投保自身 0.00046 ETH 手续费", "换低 gas 的 EVM 链 / L2（BOT Chain 968 已部署）"],
    ], [Inches(2.2), Inches(4.5), Inches(4.83)], row_h=Inches(0.6), body_size=9.5)
    notes(s, "这一页是留给我们自己挨问的：六条边界，每条都写清楚「现在是什么样」和「怎么处置」。"
             "我们不打算把任何一条说成已经解决了。")
    return s


def build(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    # 前 16 页 = 主讲（6 分钟版逐页讲；5 分钟 / 3 分钟版见 提交材料/路演PPT-讲稿与跳页表.md）
    # 后 10 页 = 备用（答辩与追问时才翻；翻页用 输入页码 + Enter）
    builders = [
        # ---- 主讲 16 页 ----
        s01_cover, s02_person, s03_conflict, s04_math, s05_what_if, s06_idea,
        s07_flow, s08_demo, s09_punchline, s10_ai1, s11_ai2, s12_ai3,
        s13_onemore, s14_compare, s15_cost, s17_close,
        # ---- 备用 10 页 ----
        s04b_problem, s04c_target, s07b_arch, s09b_evidence, s14b_tiers,
        s17b_business, s16_spec, s18a_roadmap, s18b_limits, s18_appendix,
    ]
    for i, fn in enumerate(builders, start=1):
        slide = fn(prs, i)
        if i > 1:
            foot(slide, i)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)
    return len(builders)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "提交材料", "雨证-路演PPT.pptx")
    n = build(out)
    print(f"已生成 {n} 页 / {os.path.getsize(out)} 字节 -> {out}")
