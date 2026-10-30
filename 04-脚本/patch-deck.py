# 把 PPT 的旧口径（v3/677/ETH/5 城/40 格/229 断言）改成 2026-10-27 的现状：
# BOT Chain 968 · v5 · 人民币分 · 湖北 118 区划全量开通。
# 用法：python patch-deck.py [deck-metrics.json]
# 文本替换按 shape_id 精确落点，保留原 run 的字号/颜色，不重建段落。
import json, shutil, sys
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
DECK = HERE.parent / "01-参赛材料" / "交付物" / "PPT" / "神麻组-雨证-路演PPT-答辩版-2026-10-28.pptx"
METRICS = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "deck-metrics.json"

m = json.loads(METRICS.read_text(encoding="utf-8"))
n = lambda x: f"{x}"

def rmb(cents):
    return f"{cents/100:.2f}".rstrip("0").rstrip(".")

hub = m["hubCount"]; region = m["regionCount"]; rain = m["rain"]; p24 = m["p24"]
pr = m["price"]; pool = f"{m['pool']/100:,.2f}"; claims = m["claims"]
p0 = m["p0"]; p24f = lambda k: f"{p24[k]/100:.2f}"

REPL = {
    # S7 参数页
    (7, 12): f"{region} 个区划码",
    (7, 13): "按区划独立定价，跨区不共享同一档线",
    (7, 32): f"{rmb(pr['min12'])}–{rmb(pr['max24'])} 元",
    (7, 33): f"{pr['cells']} 格费率表，按区划 × 时长逐格算，不手填",
    # S8 三条链上证据
    (8, 248): "PAYOUT CAP",
    (8, 249): "100",
    (8, 250): "元",
    (8, 251): "触发到账是分钟级。判定与赔付全程上链留痕；链上成本按 BOT 计价，批量摊薄后每单不足一元 —— 传统保险的钱花在「查」，我们的钱花在「给」。",
    (8, 256): f"买入 {p0['base']} mm，判定 {p0['judge']} mm；增量 0 mm",
    (8, 261): f"湖北 {hub} 区划 · 逐条开通",
    (8, 262): f"{hub} 条注册事件 + {hub} 条首笔喂价，全部可核验",
    (8, 263): n(hub),
    (8, 264): "已开通",
    (8, 267): "全池快照 · 2026-10-27",
    (8, 268): f"已收保费 {pool} 元，累计赔付 {claims} 笔",
    (8, 269): n(claims),
    (8, 270): "赔付",
    # S9 现场
    (9, 104): "LIVE DEMO · BOT CHAIN 968 · v5",
    (9, 123): f"五城雨量直读链上：上海 {rain['310100']}、武汉 {rain['420100']}、北京 {rain['110100']}、广州 {rain['440100']}、成都 {rain['510100']} mm；24h 档线 50 / 100，成都未到第一档。",
    (9, 124): f"判定与赔付逐条留痕：968 上 {m['policyCount']} 张保单、{claims} 笔赔付，{region} 个区划每次喂价都可核验。每条右侧都能点「核验」。",
    (9, 125): f"武汉 24 小时：保费 {rmb(p24['420100'])} 元（{p24['420100']} 分）、档线 50 / 100 mm —— 三个数都是合约函数返回的。",
    (9, 126): "三个面板读的是同一份合约。纯静态 HTML，没有后端、没有数据库 —— 浏览器直接向 968 要数：premiumOf(regionId, hours, 0, 0, 1)、entryThresholdOf(hours)，都不是页面算的。",
    # S10 精算
    (10, 284): "11 YEARS · 118 REGIONS",
    (10, 287): "PRICE SPREAD",
    (10, 288): f"{pr['ratio']}×",
    (10, 289): f"湖北 24 小时档最贵 {rmb(pr['max24'])} 元（汉南区）、最便宜 {rmb(pr['min24'])} 元（房县），相差 {pr['ratio']} 倍 —— 用一个价卖全省：贵的城市赔穿，便宜的城市被逆选择。",
    (10, 293): f"{p24f('110100')}~{p24f('310100')} 元",
    (10, 295): f"{p24f('440100')} 元",
    (10, 297): f"{p24f('510100')} 元",
    (10, 299): f"{p24f('420100')} 元",
    # S11 架构
    (11, 154): "updateRainfallBatch · rejectFeed · setUnderwriting · submitJudgment · claim · fundPool",
    (11, 155): f"Solidity 合约：v5 在 BOT Chain 968 实链运行，{region} 个区划全部开通",
    (11, 161): "24 组自动化门禁，含「每条留痕的哈希都能当场重算」",
    (11, 166): "index-v5.html 投保与看板 / rider.html 保单 / console.html 台账 / verify-v5.html 核验",
    # S16 对比
    (16, 347): "链上降雨量是否越过合约档线（v5 · 24h 档 50 / 100 mm）",
    # S18 团队
    (18, 13): f"规则设计、{pr['cells']} 格费率表、历史回测；答辩主讲",
    (18, 18): "合约 v1–v5、keeper 与门禁脚本、留痕体系",
    (18, 27): "合约 v1–v5，Sepolia 与 BOT Chain 968 跑通",
    (18, 29): "4 个 v5 页面 + 单文件核验台，纯静态、零后端、只读链",
    (18, 31): "24 组自动化门禁，含每条哈希当场重算",
    (18, 33): "11 年逐小时回测：5 城各 96,432 小时；湖北 118 区划全量定价",
    (18, 37): "从取数、开通、定价到判定的闭环已在 968 跑通",
    # S19 尾页
    (19, 38): "github.com/juice-butterfly/rainproof\u3000·\u3000演示站默认读 BOT Chain 968 · v5",
}

def set_text(shape, text):
    tf = shape.text_frame
    par = tf.paragraphs[0]
    runs = par.runs
    if not runs:
        par.add_run().text = text if hasattr(par, "add_run") else text
        if not runs:  # older python-pptx
            r = par.add_run(); r.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""

prs = Presentation(DECK)
hit = 0
for slide_no, shape_id, text in [(s, i, t) for (s, i), t in REPL.items()]:
    shape = next((sh for sh in prs.slides[slide_no - 1].shapes if sh.shape_id == shape_id), None)
    if shape is None:
        print(f"MISS slide {slide_no} shape {shape_id}")
        continue
    set_text(shape, text)
    hit += 1

# S10 三 run 的小标题：只换中间那段
hdr = next(sh for sh in prs.slides[9].shapes if sh.shape_id == 291)
hdr.text_frame.paragraphs[0].runs[1].text = "968 v5 · 24h 档"

bak = DECK.with_suffix(".pptx.bak-2026-10-27")
if not bak.exists():
    shutil.copy2(DECK, bak)
# 原文件可能正被 PowerPoint 打开（~$ 锁）：写不进去就落一份同目录副本，绝不丢改动
try:
    prs.save(DECK)
    out = DECK
    print(f"patched {hit} shapes -> {DECK.name}; backup {bak.name}")
except PermissionError:
    out = DECK.with_name(DECK.stem + "-968v5.pptx")
    prs.save(out)
    print(f"原文件被占用，已另存为 {out.name}（请关闭 PowerPoint 后替换回 {DECK.name}）；backup {bak.name}")

# 读回自检
chk = Presentation(out)
bad = []
for (s, i), t in REPL.items():
    sh = next((x for x in chk.slides[s - 1].shapes if x.shape_id == i), None)
    got = sh.text_frame.text if sh is not None else "<missing>"
    if got != t:
        bad.append((s, i, got, t))
print("verify:", "OK" if not bad else bad)
