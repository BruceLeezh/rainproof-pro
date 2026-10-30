import sys, re
from pptx import Presentation

p = r"C:\Users\lenovo\Desktop\汉客松\rainproof\提交材料\神麻组-雨证-路演PPT1.pptx"
prs = Presentation(p)
print("slides:", len(prs.slides))
pat = re.compile(r"(677|968|Sepolia|0\.0\d+|CHAIN_ID|主网|测试网)")
for i, s in enumerate(prs.slides, 1):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip().replace("\n", " / ")
            if t:
                texts.append(t)
        if sh.has_table:
            for r in sh.table.rows:
                texts.append(" | ".join(c.text.strip() for c in r.cells))
    joined = " ⟂ ".join(texts)
    hits = [x for x in pat.findall(joined)]
    flag = "★" if hits else " "
    print(f"{flag} slide{i:02d}: {joined[:400]}")
