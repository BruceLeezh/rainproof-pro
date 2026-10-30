# -*- coding: utf-8 -*-
"""看一份 pptx 有多少页、每页文本、是否英文（用于判断 B3 的「PPT 英文版」是否已存在）。"""
import sys

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
pr = Presentation(sys.argv[1])
print("slides =", len(pr.slides))
ascii_total = nonascii_total = 0
for i, s in enumerate(pr.slides, 1):
    parts = []
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip().replace("\n", " | ")
            if t:
                parts.append(t)
    line = " || ".join(parts)
    for ch in line:
        if ord(ch) < 128:
            ascii_total += 1
        else:
            nonascii_total += 1
    print(f"[{i:02d}] {line[:180]}")
tot = ascii_total + nonascii_total
print(f"字符构成：ASCII {ascii_total} / 非 ASCII {nonascii_total} = 英文占比 {ascii_total / max(1, tot):.1%}")
