import sys, zipfile
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
p = r"C:\Users\lenovo\Desktop\汉客松\rainproof\提交材料\神麻组-雨证-路演PPT1.pptx"
pr = Presentation(p)
print("页数 =", len(pr.slides), "| 尺寸 =", pr.slide_width, pr.slide_height)
z = zipfile.ZipFile(p)
media = [n for n in z.namelist() if n.startswith("ppt/media/")]
print("媒体部件:", [(n.split("/")[-1], z.getinfo(n).file_size) for n in media][:12])
for i, s in enumerate(pr.slides, 1):
    txt = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txt.append(sh.text_frame.text.strip().split("\n")[0])
    print("%2d | %s" % (i, " · ".join(txt[:3])[:110]))
