import sys
from pptx import Presentation
from pptx.util import Emu

p = sys.argv[1]
prs = Presentation(p)
print("slides:", len(prs.slides), "size:", prs.slide_width, prs.slide_height)
for i, s in enumerate(prs.slides, 1):
    print(f"\n===== slide {i} =====")
    for sh in s.shapes:
        if sh.has_text_frame:
            t = "\n".join(r.text for para in sh.text_frame.paragraphs for r in para.runs).strip()
            if t:
                print(t)
        if sh.has_table:
            for row in sh.table.rows:
                print(" | ".join(c.text for c in row.cells))
