# -*- coding: utf-8 -*-
"""26 → 29 页：补社会价值、附录 A 技术证据、附录 B 参考来源。"""
import os, sys, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from dark_lib import T, CARD, R, RULE, kicker, notes, CYAN, BLUE, AMBER, WHITE, MONO, YAHEI

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "神麻组-雨证-路演PPT-答辩版-2026-10-28.pptx")
X0, CW, TOTAL = 1.00, 11.55, 29
prs = Presentation(DECK)
assert len(prs.slides) == 26, len(prs.slides)


def dup(idx=12):
    src = prs.slides[idx]
    sl = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(sl.shapes):
        ph._element.getparent().remove(ph._element)
    for i, shp in enumerate(src.shapes):
        if i > 1:
            break
        sl.shapes._spTree.append(copy.deepcopy(shp._element))
    return sl


def head(sl, title, en, chapter, kick):
    T(sl, X0, 0.395, 10.6, 0.52, title, size=27, bold=True, wrap=False, ls=1.12)
    T(sl, X0, 1.045, 9.4, 0.24, en, size=12, font=MONO, wrap=False, alpha=52)
    T(sl, 9.30, 0.615, 3.03, 0.22, chapter, size=12, font=MONO, spc=150,
      align=PP_ALIGN.RIGHT, wrap=False, alpha=34)
    kicker(sl, X0, 1.25, kick, color=CYAN, size=10.5, w=CW)


def foot(sl, no):
    T(sl, 11.10, 6.99, 1.22, 0.20, f"{no:02d} / {TOTAL}", size=11.25, font=MONO,
      align=PP_ALIGN.RIGHT, wrap=False, alpha=40)


made = []

# ============================== 社会价值与品类延展 ==============================
sl = dup()
head(sl, "补上的是一个缺口，不是一个卖点", "Filling a gap, not selling a feature", "05 / IMPACT",
     "IMPACT · 为什么这件事值得做，以及它能长到哪一步")
CARDS = [
    ("01", "补上保障缺口", "700 万", "两大平台约 700 万量级即配骑手，「收入中断」这一块没有产品覆盖 —— 现有保障只保意外伤害，不保没收入。小额、高频、可负担，是普惠保险的典型形态。"),
    ("02", "踩上政策节奏", "2026", "人社部等九部门《关于扩大新就业形态人员职业伤害保障试点工作的通知》（人社部发〔2025〕24 号）把即时配送纳入扩大范围，2026 年推广全国。天气收入保障与职业伤害保障互补成一套。"),
    ("03", "同一引擎换参数", "3 类", "高温津贴触发 · 台风停工保障 · 空气质量户外作业保障 —— 从骑手延展到物流、建筑、环卫等户外劳动者；城市级天气风险定价能力可直接输出给险企。"),
]
w = (CW - 0.22 * 2) / 3
for i, (n, lab, val, sub) in enumerate(CARDS):
    x = X0 + i * (w + 0.22)
    CARD(sl, x, 1.60, w, 2.62)
    T(sl, x + 0.22, 1.74, w - 0.44, 0.18, n, size=8.5, font=MONO, color=CYAN, spc=150, wrap=False)
    T(sl, x + 0.22, 1.95, w - 0.44, 0.22, lab, size=11.5, bold=True, wrap=False)
    T(sl, x + 0.22, 2.22, w - 0.44, 0.44, val, size=26, bold=True,
      color=AMBER if i == 1 else CYAN, wrap=False)
    T(sl, x + 0.22, 2.74, w - 0.46, 1.36, sub, size=9.6, color=WHITE, alpha=56, ls=1.38)

RULE(sl, X0, 4.42, CW, a=8)
T(sl, X0, 4.56, CW, 0.24, "第二条曲线 · 为什么精算引擎必须可复算", size=10, font=MONO,
  color=CYAN, spc=150, wrap=False)
T(sl, X0, 4.86, CW, 0.80,
  "同一套「城市级天气 → 触发概率 → 费率」的引擎，换一个灾种、换一个人群就能出一张新费率表。"
  "但它能被输出给持牌险企的前提是：对方能自己跑一遍、拿到同一个数。所以可复算不是技术洁癖，是商业模式的前提。",
  size=10.8, color=WHITE, alpha=68, ls=1.45)
T(sl, X0, 5.86, CW, 0.60, [
    ("一句话　", {"color": AMBER, "size": 11.5, "bold": True}),
    ("我们卖的不是一份保单，是一套「可被独立复核的天气风险定价能力」。", {"color": WHITE, "size": 11.5, "alpha": 72}),
])
notes(sl, "这一页讲格局：不吹市场规模，讲这个能力还能用在哪。政策那条线只陈述文件事实，不引申政策背书。")
made.append((sl, 21))

# ============================== 附录 A · 技术证据清单 ==============================
sl = dup()
head(sl, "附录 A · 这些不是「可提供」，是现在就能跑", "Appendix A — runnable evidence, not promises", "07 / APPENDIX A",
     "TECHNICAL EVIDENCE · 技术评委可以逐条深挖")
ROWS = [
    ("自动化门禁", "28 组入口 · 596 项断言全绿", "cd 07-测试工具 && npm test", "现场可复跑"),
    ("精算引擎自检", "509 项逐项断言（11 年数据口径）", "随 npm test 一并执行", "可复跑"),
    ("V6 部署核验", "12 项：源码 / ABI / 部署字节码 / 运行时逐字节", "node 07-测试工具/check-v6-deployment.js --onchain", "链上比对"),
    ("V6 隔离链", "23 项：批量配置 · 价格守卫 · 承保容量 · 来源数 · 7 天宽限 · 重复领取", "e2e:v6-review", "本地链"),
    ("历史回放", "2024-06 武汉（三天 230mm）· 2026-08 上海，投保→验真→判定→赔付全闭环", "07-测试工具/check-replay.js", "可复算"),
    ("判不赔记录", "上海 12h 窗口增量 0mm（档线 30mm）→ 拒赔并留痕", "verify-v6 第 7 条", "可复算"),
]
y = 1.58
for name, what, how, tag in ROWS:
    CARD(sl, X0, y, CW, 0.72)
    R(sl, X0, y + 0.14, 0.028, 0.44, fill=CYAN, alpha=100)
    T(sl, X0 + 0.24, y + 0.11, 1.95, 0.24, name, size=11.5, bold=True, wrap=False)
    T(sl, X0 + 2.24, y + 0.12, 5.55, 0.24, what, size=9.8, color=WHITE, alpha=64, wrap=False)
    T(sl, X0 + 7.95, y + 0.13, 2.95, 0.22, how, size=8.8, font=MONO, color=WHITE, alpha=42, wrap=False)
    T(sl, 11.95, y + 0.11, 0.75, 0.24, tag, size=9.2, color=AMBER, align=PP_ALIGN.RIGHT, wrap=False)
    y += 0.80
T(sl, X0, 6.42, CW, 0.40,
  "链上可查：合约 0xcDe1Bb9b8d61E55e194f18B6836fF1DBd49bC61F · 部署区块 28,460,066 · 运行时 22,012 字节 · "
  "运行时哈希 0xeb7f…4100 · 2/3 多签 0x21C8…0d98", size=9.2, font=MONO, color=WHITE, alpha=50, ls=1.35)
notes(sl, "附录不放正文，被问「怎么证明」时直接翻。重点是「现在就能跑」，不是「可以提供」。")
made.append((sl, 27))

# ============================== 附录 B · 参考与数据来源 ==============================
sl = dup()
head(sl, "附录 B · 每一个外部数字都有出处", "Appendix B — every external figure is sourced", "07 / APPENDIX B",
     "REFERENCES · 访问日期与原始链接登记在《引用与待核清单》")
GROUPS = [
    ("骑手规模与收入", [
        "美团研究院 2024 年数据（工人日报 2025-05-28 报道）",
        "蜂鸟活跃骑手 ≈400 万（亿欧 2025-08）",
        "高频骑手月收入 7,230–10,100 元（2025 Q1）",
        "一线熟练骑手 11,547 元/月（2024 全年）",
    ]),
    ("国家标准与气象", [
        "GB/T 28592—2012《降水量等级》表 1（12h / 24h 档线）",
        "ERA5 再分析逐小时降水（2015-10 → 2026-09）",
        "Open-Meteo 聚合接口（ECMWF / GFS / ICON）",
        "气象灾害预警信号（暴雨蓝黄橙红）",
    ]),
    ("监管与政策", [
        "《保险法》第六条（保险业务由保险公司经营）",
        "国发〔2014〕29 号（探索天气指数保险）",
        "《气象信息服务管理办法》（气象局令第 27 号）第七条",
        "人社部发〔2021〕56 号、人社部发〔2025〕24 号",
    ]),
    ("竞品与赛道", [
        "广东巨灾指数保险（2016 首创，10 市试点）公开报道",
        "Arbol / Raincoat / FloodFlash 官方站点",
        "美团 2025「夏季清凉行动」防暑险（平台付费习惯）",
        "平台骑手规模 300 万为美团 2025-06-26 官方口径",
    ]),
]
colw = (CW - 0.30) / 2
for gi, (gl, items) in enumerate(GROUPS):
    cx = X0 + (gi % 2) * (colw + 0.30)
    cy = 1.58 + (gi // 2) * 2.42
    CARD(sl, cx, cy, colw, 2.22)
    T(sl, cx + 0.24, cy + 0.16, colw - 0.48, 0.22, gl, size=9.5, font=MONO, color=CYAN, spc=150, wrap=False)
    yy = cy + 0.48
    for it in items:
        T(sl, cx + 0.24, yy, 0.14, 0.22, "·", size=12, color=CYAN, wrap=False)
        T(sl, cx + 0.44, yy + 0.01, colw - 0.70, 0.36, it, size=9.4, color=WHITE, alpha=62, ls=1.32)
        yy += 0.42
T(sl, X0, 6.46, CW, 0.34,
  "核实纪律：全部外部引用已联网核实并逐条登记于《引用与待核清单》，标注访问日期与原始链接；"
  "无法核实的数字一律不用，缺口写「待补」而不是估算。", size=9.4, color=WHITE, alpha=52, ls=1.35)
notes(sl, "附录 B 的作用是先手：评委质疑某个数字，直接翻到这页给来源；没有来源的我们就不写。")
made.append((sl, 28))

# ============================== 重排 ==============================
sldIdLst = prs.slides._sldIdLst
elements = list(sldIdLst)
assert len(elements) == 26 + len(made)
new_els, old_els = elements[26:], elements[:26]
result = list(old_els)
for el, (_, target) in zip(new_els, made):
    result.insert(target, el)
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in result:
    sldIdLst.append(el)

# ============================== 页码 / 结构图 / 封面 ==============================
for i, s in enumerate(prs.slides):
    found = False
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if len(t) <= 12 and "/" in t and t.replace("/", "").replace(" ", "").isdigit():
            found = True
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = f"{i+1:02d} / {TOTAL}"
    if not found:
        foot(s, i + 1)

def rewrite_col(sl, pred, vals):
    """按 y 从上到下依次赋值，避免「8%」这类重复值被一起改掉"""
    cells = []
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t and pred(t):
            cells.append(sh)
    cells.sort(key=lambda s: (s.top, s.left))
    for sh, v in zip(cells, vals):
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.text = v
    return len(cells)


sl1 = prs.slides[1]
n1 = rewrite_col(sl1, lambda t: t.startswith("P0") or t.startswith("P1") or t.startswith("P2"),
                 ["P01–P02", "P03–P05", "P06–P08", "P09 + 实跑", "P10–P13", "P14–P22", "P23–P29"])
n2 = rewrite_col(sl1, lambda t: t.startswith("共"),
                 ["共 2 页", "共 3 页", "共 3 页", "共 1 页 + 实跑", "共 4 页", "共 9 页", "共 7 页"])
n3 = rewrite_col(sl1, lambda t: t.endswith("%") and len(t) <= 4,
                 ["8%", "11%", "11%", "16%", "14%", "26%", "14%"])
print("结构页列改写：", n1, n2, n3)

# 权重条重画（8/11/11/16/14/26/14）
BARW = {"Rectangle 13": 8, "Rectangle 21": 11, "Rectangle 29": 11, "Rectangle 37": 16,
        "Rectangle 45": 14, "Rectangle 53": 26, "Rectangle 61": 14}
for sh in prs.slides[1].shapes:
    if sh.name in BARW:
        sh.width = Emu(int(Pt(BARW[sh.name] * 2.6)))

for sh in prs.slides[1].shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "26 PAGES" in r.text:
                    r.text = r.text.replace("26 PAGES", "29 PAGES")
                if r.text.strip() == "团队与收束":
                    r.text = "团队 · 收束 · 附录"
for sh in prs.slides[0].shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "26 页" in r.text:
                    r.text = r.text.replace("26 页", "29 页")

prs.save(DECK)
print("完成，共", len(prs.slides), "页")
