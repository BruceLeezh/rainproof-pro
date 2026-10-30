# -*- coding: utf-8 -*-
"""把答辩版扩成完整版（26 页）：补齐评审维度里缺的 7 页，并统一重编页码与结构图。

新增页（沿用 dark_lib 同一套 Token：BG #0A1422 / CYAN #22D3EE / BLUE #2D7FF9 / AMBER #F5A524，
微软雅黑 + Consolas，白 4% 圆角卡片，位置与既有页同一栅格）：
  P11 精算基础        —— 11 年逐小时数据，专业深度
  P15 竞争格局        —— 四类玩家与空白位，可行性
  P16 护城河与弱点    —— 诚实版，可行性/团队
  P18 单位经济        —— 每一元保费的去向，商业价值
  P19 财务预测        —— 三年三情形，商业价值
  P21 风险登记册      —— 风险变参数，专业深度
  P25 资源需求        —— Ask，创业赛必需

素材全部来自 01-项目计划书-草稿.md / 30 页计划书 PPT / V6 验收文档，数字未经改写。
"""
import os, sys, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from dark_lib import (T, CARD, R, RULE, kicker, retitle, notes,
                      CYAN, BLUE, AMBER, WHITE, YAHEI, MONO, BG)

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "神麻组-雨证-路演PPT-答辩版-2026-10-28.pptx")

X0, CW = 1.00, 11.55          # 内容区左边界与宽度（英寸）
FINAL_TOTAL = 26

prs = Presentation(DECK)
# 上一轮脚本可能已跑过；保证只跑一次
assert len(prs.slides) == 19, f"期望 19 页，实得 {len(prs.slides)}"


# ------------------------------------------------------------------ 基础构件
def dup(sl_idx):
    """复制一页的骨架（底色 / 章节竖条），标题与正文一律由调用方按同一栅格重画。"""
    src = prs.slides[sl_idx]
    sl = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(sl.shapes):
        ph._element.getparent().remove(ph._element)
    for i, shp in enumerate(src.shapes):
        if i > 1:                      # 0 底色 · 1 标题前竖条
            break
        sl.shapes._spTree.append(copy.deepcopy(shp._element))
    return sl


def head(sl, title, en, chapter, kick):
    T(sl, X0, 0.395, 10.6, 0.52, title, size=27, bold=True, wrap=False, ls=1.12)
    T(sl, X0, 1.045, 9.4, 0.24, en, size=12, font=MONO, wrap=False, alpha=52)
    T(sl, 9.30, 0.615, 3.03, 0.22, chapter, size=12, font=MONO, spc=150,
      align=PP_ALIGN.RIGHT, wrap=False, alpha=34)
    kicker(sl, X0, 1.25, kick, color=CYAN, size=10.5, w=CW)


def foot(sl, no):
    T(sl, 11.10, 6.99, 1.22, 0.20, f"{no:02d} / {FINAL_TOTAL}", size=11.25,
      font=MONO, align=PP_ALIGN.RIGHT, wrap=False, alpha=40)


def note_line(sl, y, text, size=9.5):
    T(sl, X0, y, CW, 0.30, text, size=size, color=WHITE, alpha=55, ls=1.35)


def card(sl, x, y, w, h, num, label, value, sub, accent=WHITE):
    CARD(sl, x, y, w, h)
    T(sl, x + 0.22, y + 0.15, w - 0.44, 0.18, num, size=8.5, font=MONO, color=CYAN, spc=150, wrap=False)
    T(sl, x + 0.22, y + 0.36, w - 0.44, 0.20, label, size=10.5, color=WHITE, alpha=60, wrap=False)
    T(sl, x + 0.22, y + 0.58, w - 0.44, 0.38, value, size=21, bold=True, color=accent, wrap=False)
    T(sl, x + 0.22, y + 1.00, w - 0.46, h - 1.06, sub, size=9.2, color=WHITE, alpha=50, ls=1.32)


def row(sl, y, label, desc, tag, tag_color=AMBER, h=0.70, label_w=3.05):
    CARD(sl, X0, y, CW, h)
    R(sl, X0, y + 0.14, 0.028, h - 0.28, fill=CYAN, alpha=100)
    T(sl, X0 + 0.24, y + 0.13, label_w, 0.24, label, size=12.5, bold=True, wrap=False)
    T(sl, X0 + 0.24 + label_w, y + 0.16, CW - label_w - 2.2, 0.24, desc, size=10.5,
      color=WHITE, alpha=62, wrap=False)
    T(sl, 12.55 - 1.95, y + 0.13, 1.75, 0.24, tag, size=11, color=tag_color,
      align=PP_ALIGN.RIGHT, wrap=False)


def grid(sl, items, y0, card_h, gap=0.20):
    """items: [(num,label,value,sub,accent)]，按宽度自动分列"""
    n = len(items)
    cols = 3 if n > 4 else (2 if n > 2 else n)
    rows = (n + cols - 1) // cols
    w = (CW - gap * (cols - 1)) / cols
    for i, it in enumerate(items):
        c, r = i % cols, i // cols
        card(sl, X0 + c * (w + gap), y0 + r * (card_h + gap), w, card_h,
             it[0], it[1], it[2], it[3], it[4] if len(it) > 4 else WHITE)


def table(sl, header, rows, y0, col_w, row_h=0.52, head_h=0.30):
    x = X0
    for txt, w in zip(header, col_w):
        T(sl, x + 0.18, y0, w - 0.30, 0.22, txt, size=9.5, font=MONO, color=CYAN,
          spc=150, wrap=False)
        x += w
    RULE(sl, X0, y0 + head_h, CW, a=10)
    y = y0 + head_h + 0.06
    for r in rows:
        CARD(sl, X0, y, CW, row_h)
        x = X0
        for j, (txt, w) in enumerate(zip(r, col_w)):
            bold = j == 0
            T(sl, x + 0.18, y + (row_h - 0.22) / 2, w - 0.28, 0.24, txt,
              size=11 if bold else 10.5, bold=bold,
              color=WHITE if j == 0 else (AMBER if j == 1 else WHITE),
              alpha=None if (j <= 1) else 72, wrap=False)
            x += w
        y += row_h + 0.08
    return y


def bullets(sl, x, y, w, title, items, gap=0.42, tcolor=CYAN):
    T(sl, x, y, w, 0.22, title, size=10, font=MONO, color=tcolor, spc=150, wrap=False)
    yy = y + 0.30
    for it in items:
        T(sl, x, yy, 0.14, 0.22, "·", size=13, color=CYAN, wrap=False)
        T(sl, x + 0.20, yy + 0.015, w - 0.20, 0.30, it, size=10.2, color=WHITE,
          alpha=66, ls=1.32)
        yy += gap
    return yy


# ================================================================== 新增页
made = []      # (slide, 目标序号 0-based)

# ---------- P11 精算基础 ----------
sl = dup(12)
head(sl, "定价是从 11 年逐小时数据里算出来的", "Priced off 11 years of hourly rainfall", "04 / ACTUARY",
     "ACTUARIAL BASIS · 不是把历史均值填进公式")
grid(sl, [
    ("01", "样本跨度", "11 年", "2015-10 → 2026-09，五城逐小时降水（ERA5 再分析），公开可复算"),
    ("02", "单城样本量", "96,432 h", "五城同口径；滑动窗口覆盖 6 / 12 / 24 / 48 / 72h"),
    ("03", "置信区间", "4,000 次", "自举重采样取 97.5% 上界作为期望赔付，不用点估计"),
    ("04", "城市极差", "3.2 倍", "24h 首档触发概率：北京约 0.4% → 广州约 1.4%，统一价必亏"),
    ("05", "逆选择量化", "5.7–10.4×", "挑最优时点投保，期望赔付/敞口中位比；条件触发率 2.6–6.7 倍"),
    ("06", "计算精度", "0.1 mm", "全程整数十分位累加，零浮点漂移，跨机器复现同一结果"),
], 1.62, 1.62)
note_line(sl, 6.02, "三条安全边际不是话术，是引擎里的常数：取 97.5% 上界（不是点估计）· 向上取整 · 赔付率红线 60%。"
                    "任何人拿同一份数据切片跑同一个引擎，得到同一张费率表 —— 这是定价可被审计的前提。", 9.8)
note_line(sl, 6.50, "刻意的取舍：只做 12h / 24h（国标没有 48h / 72h，线性外推会造成「多付 4 倍保费、中奖率降一半」）；"
                    "不设特大暴雨档（11 年五城 24h ≥ 250mm 零命中，写进条款却永远赔不到的档是虚假承诺）。", 9.2)
notes(sl, "讲这一页只用一句话：费率不是我拍的，是 11 年逐小时数据算出来的，而且换个人跑同一个引擎得同一个数。\n"
          "如果评委追问极端值：24h 档 VaR99 为负，准备金下限取 0；真正吃紧的是同城集中赔付，靠跨城市分散（10 对城市仅 2 对显著相关）化解。")
made.append((sl, 10))

# ---------- P15 竞争格局 ----------
sl = dup(12)
head(sl, "四类玩家，空白在个人小额高频这一格", "Four kinds of players, one empty square", "05 / COMPETITION",
     "COMPETITION · 谁已经在做，谁还没做到这一层")
row(sl, 1.60, "传统骑手意外险", "平台合作的意外险 / 第三者责任险：不保收入中断，理赔仍需报案定损", "不覆盖", WHITE, 0.74, 3.15)
row(sl, 2.44, "平台自保障", "恶劣天气补贴 / 关怀金：无精算、无定价、无分保，成本中心，不可审计", "不是保险", WHITE, 0.74, 3.15)
row(sl, 3.28, "国内参数化先例", "广东巨灾指数保险（2016 首创，10 市试点）、农业天气指数险：证明参数触发被监管接受，但只到城市级巨灾", "无个人场景", WHITE, 0.74, 3.15)
row(sl, 4.12, "国际参数化玩家", "Arbol / Raincoat / FloodFlash：佐证赛道成熟；中国即配个人场景的覆盖仍待验证", "未进入", WHITE, 0.74, 3.15)
RULE(sl, X0, 5.06, CW, a=8)
T(sl, X0, 5.20, CW, 0.60, [
    ("我们的位置　", {"color": AMBER, "size": 12.5, "bold": True}),
    ("把参数化从「城市巨灾」压到「一个骑手的一天」，再用链上哈希把「可验证」做成产品的一部分 —— "
     "这一格现在没有人占，因为占它需要同时具备精算数据、验真工程与可复算的存证。", {"color": WHITE, "size": 11.5, "alpha": 70}),
], ls=1.4)
notes(sl, "这一页回答「为什么是你们」。不要说「我们没有竞争对手」，要说清楚四类玩家各自做到了哪一层、"
          "为什么跨不到我们这一层：城市级巨灾不解决个人小额高频，平台补贴不是保险、不可审计。")
made.append((sl, 14))

# ---------- P16 护城河与弱点 ----------
sl = dup(12)
head(sl, "四条护城河，三个弱点，都写在这一页", "Four moats, three weaknesses, all disclosed", "05 / MOAT",
     "MOAT & WEAKNESS · 只讲能验证的部分")
bullets(sl, X0, 1.62, 5.55, "四条护城河", [
    "数据与精算：244 格费率的完整复算链，换个人跑同一引擎得同一个数",
    "验真层：三模型交叉验证的工程实现，不是「买了 API 就算有」",
    "可验证理赔：规则与证据全程哈希存证，第三方可复算同一结论",
    "先发：已建成、已回测、已真实赔付过，并已完成 V6 迁移",
], gap=0.62, tcolor=CYAN)
bullets(sl, 6.95, 1.62, 5.60, "三个弱点与应对", [
    "团队小 → 按规则组队 2–4 人 + 指导教师 1–2 名，跨学科配置",
    "尚无持牌伙伴 → 输出型定位（不做承保主体）降低门槛，赛前争取 1 家意向",
    "预言机单点 → 三源交叉已实现，国内合规数据源切换路径已明确",
], gap=0.62, tcolor=AMBER)
RULE(sl, X0, 4.30, CW, a=8)
T(sl, X0, 4.44, CW, 0.24, "还没解决、也不假装解决的三件事", size=10, font=MONO, color=CYAN, spc=150, wrap=False)
for i, (h_, d_) in enumerate([
    ("基差风险", "雨量达标但个体未受损 —— 二期接平台订单数据做「实际停单」触发，条款里先明示参数化属性"),
    ("数据源依赖", "三模型 + 容差 + 离群拒收留痕；条款里先约定断供兜底口径"),
    ("身份核验", "保单绑定骑手地址、由平台批量签发，单骑手无法自证 —— 靠 B2B2C 绕开"),
]):
    y = 4.76 + i * 0.56
    T(sl, X0, y, 1.55, 0.24, h_, size=10.5, bold=True, color=AMBER, wrap=False)
    T(sl, X0 + 1.70, y, CW - 1.70, 0.40, d_, size=10, color=WHITE, alpha=62, ls=1.3)
notes(sl, "诚实版护城河。主动报弱点，比等评委问出来强。三个弱点都要接一句「已经在做什么」。")
made.append((sl, 15))

# ---------- P18 单位经济 ----------
sl = dup(12)
head(sl, "每一元保费去哪了", "Where every yuan of premium goes", "05 / UNIT ECONOMICS",
     "UNIT ECONOMICS · 以广州 24h 平台团体价 2.2 元为例")
grid(sl, [
    ("01", "期望赔付", "34%", "0.742 ÷ 2.2；按 97.5% 上界口径，不是点估计", AMBER),
    ("02", "固定成本", "14%", "0.30 ÷ 2.2；出单 / 结算 / 存证的边际成本", WHITE),
    ("03", "再保 + 风险准备", "15%", "只在雨季前前置；24h 档 VaR99 为负，准备金下限取 0", WHITE),
    ("04", "渠道与运营", "25%", "B2B2C：平台通道费 + 服务费", WHITE),
    ("05", "科技毛利", "12%", "科技服务费口径，不碰承保利润", CYAN),
], 1.62, 1.55)
note_line(sl, 4.68, "为什么这个结构能成立：期望赔付压在 34%，是因为定价取上界而非点估计、且只做 12h / 24h 两档；"
                    "固定成本压到 14%，是因为定损环节被删掉、判定由规则与外部任务完成。", 9.8)
RULE(sl, X0, 5.28, CW, a=8)
T(sl, X0, 5.42, CW, 0.24, "技术成本实测（这台机器就能跑）", size=10, font=MONO, color=CYAN, spc=150, wrap=False)
for i, (k, v) in enumerate([("机器", "2 核 2G / 50G 单机"), ("喂价一轮", "122 区划 · 2.14 s CPU（2.2% 单核）"),
                            ("峰值内存", "87 MB"), ("部署体积", "静态站 + 合约 + 脚本 19 MB")]):
    x = X0 + i * 2.90
    T(sl, x, 5.74, 1.10, 0.22, k, size=9.5, color=WHITE, alpha=48, wrap=False)
    T(sl, x + 1.10, 5.74, 1.75, 0.22, v, size=9.5, color=WHITE, alpha=78, wrap=False)
note_line(sl, 6.14, "结论：机器不是瓶颈，天气 API 的额度才是。这也是我们敢把「规模」放在第三年、而不是第一年的原因。", 9.8)
notes(sl, "单位经济这一页最好用一张嘴算完：2.2 元里 34% 赔出去、14% 是固定成本、25% 是渠道、12% 是科技毛利。"
          "评委常问「你怎么保证不赔穿」——答案是三条安全边际，不是靠涨价。")
made.append((sl, 17))

# ---------- P19 财务预测 ----------
sl = dup(12)
head(sl, "三年三情形，假设全部摊开", "Three years, every assumption disclosed", "05 / FINANCIALS",
     "FINANCIALS · 单位：万元 · 假设可替换，结论跟着变")
table(sl,
      ["指标（万元）", "Y1 试点", "Y2 扩展", "Y3 规模"],
      [["覆盖骑手", "2 万", "20 万", "80 万"],
       ["保费规模", "144", "2,080", "11,200"],
       ["赔付（自留口径）", "≈75", "≈1,080", "≈5,800"],
       ["科技服务费收入", "18–22", "250–310", "1,350–1,680"],
       ["运营成本", "≈80", "≈350", "≈900"],
       ["净利", "负（投入期）", "−40 ~ −100", "450–780"]],
      1.62, [4.35, 2.40, 2.40, 2.40], row_h=0.50)
note_line(sl, 5.42, "假设全部显性：单价按平台批发 1.2–1.4 元；投保频次随渗透从 60 天/年升至 100 天/年；"
                    "自留赔付率 52%（点估计 47% 自留 + 5% 转分保后的净额）。", 9.8)
note_line(sl, 5.82, "盈亏平衡约 25–30 万覆盖骑手，Y3 初达成；保守 / 乐观区间 ±40%。"
                    "不承诺承保利润 —— 收入结构是科技服务费（按保费比例）+ 系统服务费（按单 / 按年）。", 9.8)
T(sl, X0, 6.28, CW, 0.60, [
    ("一句实话　", {"color": AMBER, "size": 11.5, "bold": True}),
    ("Y1 是亏的，而且我们把它写出来了。这一年的目标是跑通单城闭环与条款报备，不是规模。", {"color": WHITE, "size": 10.5, "alpha": 62}),
])
notes(sl, "财务这页的关键不是数字漂亮，是假设齐。被问「凭什么 80 万骑手」就答：单城在保规模上限由平台单量决定，"
          "不由我们的销售能力决定 —— 这是 B2B2C 的特征。")
made.append((sl, 18))

# ---------- P21 风险登记册 ----------
sl = dup(12)
head(sl, "把风险清单变成定价参数", "Every risk becomes a pricing parameter", "05 / RISK REGISTER",
     "RISK REGISTER · 左边是风险，右边是它落在哪个参数上")
row(sl, 1.58, "基差风险", "雨量达标但人未受损 → 条款明示参数化属性；二期接订单数据做实际停单触发", "已量化", CYAN, 0.66, 2.55)
row(sl, 2.32, "逆选择", "看预报挑雨天买 → 量化进价格：众包自助比平台团体贵 3–6 倍", "已产品化", CYAN, 0.66, 2.55)
row(sl, 3.06, "同城集中赔付", "准备金 ≥ 在保份数 × VaR99；跨城市分散（10 对城市仅 2 对显著相关）", "已规则化", CYAN, 0.66, 2.55)
row(sl, 3.80, "数据源污染 / 故障", "三机构交叉验真 + 容差 + 离群拒收留痕；断供时暂停投保，已售保单照约履行", "已实现", CYAN, 0.66, 2.55)
row(sl, 4.54, "极端年份尾部", "台风季前分出超赔，广深沿海优先", "二期落地", AMBER, 0.66, 2.55)
row(sl, 5.28, "运营方单点", "管理操作需 2/3 多签；独立托管仍待落实 —— 这一条我们没做到位", "部分实现", AMBER, 0.66, 2.55)
note_line(sl, 6.16, "同一份清单也出现在项目计划书与代码审计里 —— 不是答辩现场临时想的。最后一条我们不粉饰：独立托管还没做，"
                    "现在靠 2/3 多签与同机密钥约束，这是已知短板。", 9.5)
notes(sl, "如果只讲一条，讲逆选择：它不是「注意一下」，而是直接进了价格表（3–6 倍差价）。"
          "最后一条主动认，比被问出来好。")
made.append((sl, 20))

# ---------- P25 资源需求 ----------
sl = dup(12)
head(sl, "我们要什么，钱花在哪", "What we need, and where it goes", "06 / ASK",
     "ASK · 三件事，先说清楚不要什么")
grid(sl, [
    ("01", "持牌合作", "1 家", "持牌险企合作意向 —— 承保主体，我们做产品设计、精算定价与自动理赔系统", CYAN),
    ("02", "渠道试点", "1 家", "平台 / 众包服务商试点意向 —— 一次决策覆盖上万骑手，边际获客成本 ≈ 0", CYAN),
    ("03", "启动资金", "150 万", "12 个月现金储备：数据与算力 30 / 团队 80 / 试点运营 40", AMBER),
], 1.62, 1.80)
RULE(sl, X0, 4.02, CW, a=8)
T(sl, X0, 4.16, CW, 0.24, "口径说明 · 150 万 ≠ 一年亏损", size=10, font=MONO, color=CYAN, spc=150, wrap=False)
T(sl, X0, 4.48, CW, 0.70,
  "Y1 运营成本 ~80 万是「当年损益口径」；150 万是「12 个月现金储备」= 80 万运营 + 70 万试点城市与合规前置投入。"
  "两者不是同一个口径，不能混读。", size=10.5, color=WHITE, alpha=66, ls=1.4)
T(sl, X0, 5.34, CW, 0.24, "奖金投向 · 先合规后落地", size=10, font=MONO, color=CYAN, spc=150, wrap=False)
T(sl, X0, 5.66, CW, 0.50,
  "优先投数据合规源切换与试点城市落地，不投市场投放 —— 这个阶段买流量没有意义，"
  "先把「数据合规 + 单城闭环 + 条款报备」三件事做实。", size=10.5, color=WHITE, alpha=66, ls=1.4)
note_line(sl, 6.36, "我们不申请承保牌照，不做资金池，不碰保费归集 —— 这是定位，也是给合作方的定心丸。", 9.5)
notes(sl, "Ask 页要短。三件事：一家持牌、一家渠道、150 万。然后立刻解释 150 万和 80 万不是一回事，避免被当成烧钱。")
made.append((sl, 24))


# ================================================================== 重排
sldIdLst = prs.slides._sldIdLst
elements = list(sldIdLst)
assert len(elements) == 19 + len(made)
new_els, old_els = elements[19:], elements[:19]
result = list(old_els)
for el, (_, target) in zip(new_els, made):     # target 升序，逐个插入即最终位次
    result.insert(target, el)
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in result:
    sldIdLst.append(el)

# ================================================================== 页码与结构图
for i, s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if len(txt) <= 12 and "/" in txt and txt.replace("/", "").replace(" ", "").isdigit():
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = f"{i+1:02d} / {FINAL_TOTAL}"

# 结构图：页面区间跟着重排（P01 索引 1）
RANGE_FIX = [("P10–P12", "P10–P13"), ("P13–P15", "P14–P21"), ("P16–P19", "P22–P26")]
for sh in prs.slides[1].shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            for a, b in RANGE_FIX:
                if r.text.strip() == a:
                    r.text = b

# 封面与结构图的时长表述
for idx, pairs in [
    (1, [("这 15 分钟怎么走", "全篇结构"), ("How the next 15 minutes are spent", "How this deck is organized"),
         ("RUN OF SHOW · 900 SECONDS", "STRUCTURE · 26 PAGES")]),
    (0, [("初赛答辩 · 15:00 · 演示 3′00″", "初赛答辩 · 全篇 26 页 · 含现场演示")]),
]:
    for sh in prs.slides[idx].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                for a, b in pairs:
                    if a in r.text:
                        r.text = r.text.replace(a, b)

prs.save(DECK)
print("已保存，共", len(prs.slides), "页")
